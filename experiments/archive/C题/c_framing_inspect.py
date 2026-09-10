"""只读核对 C 题附件结构，生成审题证据；不进行清洗或策略求解。"""

from datetime import datetime, time, timedelta
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "C题"
OUTPUT = ROOT / "experiments" / "c_framing_evidence.json"


def digest(path):
    return sha256(path.read_bytes()).hexdigest()


def minute(value):
    if isinstance(value, time):
        return value.hour * 60 + value.minute
    value = str(value)
    clock, *day = value.split("+")
    h, m = map(int, clock.split(":"))
    return h * 60 + m + (int(day[0]) * 1440 if day else 0)


def matrix_summary(rows):
    values = [v for row in rows for v in row]
    nums = [v for v in values if isinstance(v, (int, float))]
    return {
        "cells": len(values),
        "blank": sum(v is None or v == "" for v in values),
        "non_numeric": sum(not isinstance(v, (int, float)) for v in values),
        "negative": sum(v < 0 for v in nums),
        "minimum": min(nums) if nums else None,
        "maximum": max(nums) if nums else None,
    }


def main():
    paths = sorted(p for p in SOURCE.rglob("*") if p.suffix in {".pdf", ".xlsx"})
    before = {p.relative_to(ROOT).as_posix(): digest(p) for p in paths}
    evidence = {"source_sha256": before, "attachments": {}, "templates": {}}
    evidence["pdf_pages"] = [p.extract_text() for p in PdfReader(SOURCE / "C题.pdf").pages]
    expected_dates = [datetime(2025, 1, 1) + timedelta(days=i) for i in range(365)]
    for number in range(1, 5):
        path = SOURCE / "附件" / f"附件{number}.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        report = {}
        for sheet in workbook:
            rows = list(sheet.values)
            item = {"rows": sheet.max_row, "columns": sheet.max_column}
            if number == 1:
                item["values"] = matrix_summary([row[1:] for row in rows[1:]])
                item["time_grid_10_to_1440"] = [minute(row[0]) for row in rows[1:]] == list(range(10, 1441, 10))
                item["column_ranges"] = {rows[0][c]: matrix_summary([[row[c]] for row in rows[1:]]) for c in range(1, 4)}
            elif number in (2, 4):
                dates = [row[0] for row in rows[1:]]
                item["dates_match_365_days"] = dates == expected_dates
                item["duplicate_dates"] = len(dates) - len(set(dates))
                item["time_grid_10_to_1440"] = [minute(v) for v in rows[0][1:]] == list(range(10, 1441, 10))
                item["values"] = matrix_summary([row[1:] for row in rows[1:]])
            else:
                item["values"] = matrix_summary([row[2:] for row in rows[1:]])
                issue_times = []
                current_date = None
                blank_dates = 0
                for row in rows[1:]:
                    if row[0] not in (None, ""):
                        current_date = datetime.strptime(str(row[0]), "%Y-%m-%d")
                    else:
                        blank_dates += 1
                    issue_times.append(current_date + timedelta(minutes=minute(row[1])))
                expected = [d + timedelta(hours=h) for d in expected_dates for h in (0, 6, 12, 18)]
                item["issues_match_365_times_4"] = issue_times == expected
                item["blank_date_labels"] = blank_dates
                item["duplicate_issues"] = len(issue_times) - len(set(issue_times))
                item["lead_headers"] = rows[0][2:]
                item["last_forecast_target"] = issue_times[-1] + timedelta(hours=24)
            report[sheet.title] = item
        workbook.close()
        evidence["attachments"][path.name] = report

    expected_output_dates = expected_dates[31:]
    for path in sorted((SOURCE / "附件" / "附件5").glob("*.xlsx")):
        workbook = load_workbook(path, read_only=True, data_only=False)
        report = {}
        for sheet in workbook:
            rows = list(sheet.values)
            item = {"rows": sheet.max_row, "columns": sheet.max_column}
            if sheet.title in ("计划购电量", "调整购电量"):
                if path.name == "result1.xlsx":
                    labels = [row[0] for row in rows[1:]]
                else:
                    labels = list(rows[0][1:145])
                    item["dates_match_feb_to_dec"] = [row[0] for row in rows[1:]] == expected_output_dates
                    item["trailing_headers"] = rows[0][145:]
                item["interval_count"] = len(labels)
                item["first_label"] = labels[0]
                item["last_label"] = labels[-1]
                item["starts_at_10_not_0"] = [minute(v.split("-")[0]) for v in labels] == list(range(10, 1441, 10))
                item["ends_at_1450"] = [minute(v.split("-")[1]) for v in labels] == list(range(20, 1451, 10))
            else:
                item["nonempty_labels"] = [[i, list(row)] for i, row in enumerate(rows, start=1) if any(v is not None for v in row)]
            report[sheet.title] = item
        workbook.close()
        evidence["templates"][path.name] = report
    after = {p.relative_to(ROOT).as_posix(): digest(p) for p in paths}
    evidence["source_hashes_unchanged"] = before == after
    if before != after:
        raise RuntimeError("检查期间源文件发生变化，请重新核对")
    OUTPUT.write_text(json.dumps(evidence, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(f"已检查 {len(paths)} 个源文件，证据保存到 {OUTPUT.relative_to(ROOT)}")
    print(f"源文件哈希一致：{evidence['source_hashes_unchanged']}")
    for name, sheets in evidence["attachments"].items():
        for title, item in sheets.items():
            print(name, title, json.dumps({k: v for k, v in item.items() if k != "lead_headers"}, ensure_ascii=False, default=str))
    for name, sheets in evidence["templates"].items():
        item = sheets["计划购电量"]
        print(name, item["first_label"], "→", item["last_label"])


if __name__ == "__main__":
    main()
